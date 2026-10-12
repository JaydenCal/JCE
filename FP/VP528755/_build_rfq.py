"""Build final RFQ .docx files for VP528755 (TCW00728 Vickers Rd PS C6B)."""
import os, re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = "/home/user/JCE/FP/VP528755"
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
GREY = RGBColor(0x55, 0x55, 0x55)

QUOTE_DUE = "12:00 pm AEST Friday 23/10/2026"
CLOSE = "5:00 pm AEST Friday 30/10/2026"
ISSUE = "12/10/2026"
VALID_TO = "21/01/2027"


def shade(cell, hexcol):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcol)
    tcPr.append(shd)


def no_split_row(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    trPr.append(el)


def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)



def set_widths(t, widths):
    t.autofit = False
    grid = t._tbl.tblGrid
    for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
        if i < len(widths):
            gc.set(qn("w:w"), str(int(widths[i] * 567)))
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)

class Doc:
    def __init__(self):
        self.d = Document()
        sec = self.d.sections[0]
        sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
        sec.left_margin = sec.right_margin = Cm(2.0)
        sec.top_margin, sec.bottom_margin = Cm(1.6), Cm(1.8)
        st = self.d.styles["Normal"]
        st.font.name = "Calibri"
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        st.font.size = Pt(10.5)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.line_spacing = 1.1
        self.sec_no = 0

    def p(self, text="", bold=False, size=None, color=None, align=None, after=None,
          italic=False, keep=False):
        para = self.d.add_paragraph()
        self._runs(para, text, bold=bold, size=size, color=color, italic=italic)
        if align:
            para.alignment = align
        if after is not None:
            para.paragraph_format.space_after = Pt(after)
        if keep:
            para.paragraph_format.keep_with_next = True
        return para

    def _runs(self, para, text, bold=False, size=None, color=None, italic=False):
        # **bold** markup
        parts = re.split(r"(\*\*[^*]+\*\*)", text)
        for part in parts:
            if not part:
                continue
            b = bold
            if part.startswith("**") and part.endswith("**"):
                part, b = part[2:-2], True
            r = para.add_run(part)
            r.bold = b
            r.italic = italic
            if size:
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = color

    def h(self, text):
        self.sec_no += 1
        para = self.p(f"{self.sec_no}.  {text}", bold=True, size=12, color=NAVY, keep=True)
        para.paragraph_format.space_before = Pt(10)
        para.paragraph_format.space_after = Pt(4)
        return para

    def sub(self, text):
        para = self.p(text, bold=True, size=10.5, color=NAVY, keep=True)
        para.paragraph_format.space_before = Pt(4)
        return para

    def bullets(self, items):
        for it in items:
            para = self.d.add_paragraph(style="List Bullet")
            self._runs(para, it)
            para.paragraph_format.space_after = Pt(2)

    def numbered(self, items):
        for i, it in enumerate(items, 1):
            para = self.d.add_paragraph()
            para.paragraph_format.left_indent = Cm(0.8)
            para.paragraph_format.first_line_indent = Cm(-0.6)
            para.paragraph_format.space_after = Pt(2)
            self._runs(para, f"{self.sec_no}.{i}  {it}")

    def table(self, header, rows, widths, bold_last=False, font=9.5):
        t = self.d.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, htxt in enumerate(header):
            c = t.rows[0].cells[i]
            c.text = ""
            self._runs(c.paragraphs[0], htxt, bold=True, size=font, color=RGBColor(0xFF, 0xFF, 0xFF))
            shade(c, "1F3A5F")
        set_repeat_header(t.rows[0])
        for ri, row in enumerate(rows):
            cells = t.add_row().cells
            for i, val in enumerate(row):
                cells[i].text = ""
                b = bold_last and ri == len(rows) - 1
                self._runs(cells[i].paragraphs[0], str(val), bold=b, size=font)
                if b:
                    shade(cells[i], "E8EEF5")
        for row in t.rows:
            no_split_row(row)
        set_widths(t, widths)
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    def letterhead(self, title, subtitle, rfq_no, to_line):
        self.p("JCE Infrastructure Pty Ltd", bold=True, size=16, color=NAVY, after=0)
        self.p("ABN 29 699 165 663   |   Phone 0477 207 027   |   jayden@jceinfrastructure.com.au",
               size=9.5, color=GREY, after=0)
        self.p("756 Quia Road, Gunnedah NSW 2380   |   Civil and hydraulic engineering", size=9.5,
               color=GREY, after=6)
        # rule
        para = self.d.add_paragraph()
        pPr = para._p.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        bot = OxmlElement("w:bottom")
        bot.set(qn("w:val"), "single"); bot.set(qn("w:sz"), "12")
        bot.set(qn("w:space"), "1"); bot.set(qn("w:color"), "1F3A5F")
        bdr.append(bot); pPr.append(bdr)
        para.paragraph_format.space_after = Pt(8)
        self.p("REQUEST FOR QUOTATION", bold=True, size=11, color=GREY, after=2)
        self.p(title, bold=True, size=15, color=NAVY, after=2)
        self.p(subtitle, size=10.5, after=8)
        t = self.d.add_table(rows=0, cols=4)
        t.style = "Table Grid"
        info = [
            ("Date", ISSUE, "RFQ no.", rfq_no),
            ("Quote due", QUOTE_DUE, "Head tender closes", CLOSE),
            ("Client", "Townsville City Council (QLD)", "Council reference", "TCW00728 (VendorPanel VP528755)"),
            ("From", "Jayden Callen, Director", "Contact", "0477 207 027, jayden@jceinfrastructure.com.au"),
        ]
        for a, b, c, d in info:
            cells = t.add_row().cells
            for i, (val, isl) in enumerate([(a, 1), (b, 0), (c, 1), (d, 0)]):
                cells[i].text = ""
                self._runs(cells[i].paragraphs[0], val, bold=bool(isl), size=9.5)
                if isl:
                    shade(cells[i], "E8EEF5")
        set_widths(t, [2.6, 5.4, 3.3, 5.7])
        self.d.add_paragraph().paragraph_format.space_after = Pt(0)
        self.p(f"**To:** {to_line}", after=6)

    def signature(self):
        para = self.p("Yours faithfully,", keep=True)
        para.paragraph_format.space_before = Pt(8)
        self.p("", keep=True, after=14)
        self.p("Jayden Callen", bold=True, keep=True, after=0)
        self.p("Director, JCE Infrastructure Pty Ltd", keep=True, after=0)
        self.p("0477 207 027   |   jayden@jceinfrastructure.com.au", after=0)

    def footer(self, text):
        sec = self.d.sections[0]
        fp = sec.footer.paragraphs[0]
        fp.text = ""
        r = fp.add_run(text + "   |   Page ")
        r.font.size = Pt(8); r.font.color.rgb = GREY
        for code in ("PAGE",):
            r2 = fp.add_run()
            f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = code
            f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "end")
            r2._r.append(f1); r2._r.append(it); r2._r.append(f2)
            r2.font.size = Pt(8); r2.font.color.rgb = GREY

    def save(self, path):
        os.makedirs(os.path.dirname(path), exist_ok=True)
        self.d.save(path)


BACKGROUND = (
    "Townsville City Council (Council) has invited quotations under RFQ TCW00728 for an independent "
    "Superintendent and construction verification services for the Vickers Road PS C6B Diversion Project. "
    "Flow from PS C6B (North Vickers Rd, Condon) will be diverted from Condon STP to Mount St John STP. "
    "AECOM is the designer. The works will be built under three Construct Only lump sum contracts on "
    "Council's modified AS 4000-1997:"
)
CONTRACT_ROWS = [
    ["Contract 1", "Element 3: Stage 1 pressure main, DN560 PE, about 4,290 m along Kern Main Drain to Dalrymple Rd, with crossings of the Townsville Ring Road and Hervey Range Road in existing envelopers", "December 2026 to January 2028 (61 weeks)"],
    ["Contract 2", "Elements 1 and 2: PS C6B structural, pipework, access road, electrical and control upgrade, and a new odour control unit designed and constructed by the Contractor to AECOM's performance specification", "July 2027 to November 2028 (74 weeks)"],
    ["Contract 3", "Element 4: Stage 2 pressure main, DN710 PE about 7,600 m (Dalrymple Rd to Everett St, crossings of Dalrymple Rd, Woolcock St and the North Coast Railway in envelopers), then DN900 PE about 3,320 m to the Mt St John STP inlet works", "March 2028 to July 2029 (74 weeks)"],
]
BACKGROUND_2 = (
    "Each construction contract has a 12-month defects liability period. The consultancy starts on "
    "11/12/2026 (anticipated award) and its date for completion is 30/06/2030. Council will pay for hours "
    "worked at the quoted rates, up to a not-to-exceed allowance per task, under its amended AS 4122-2010 "
    "consultancy contract. Council scores the nominated individuals' capability (50%), methodology (35%), "
    "local industry participation (10%) and social benefit and Indigenous engagement (5%)."
)
JCE_INTENT = (
    "JCE Infrastructure Pty Ltd (JCE) will quote to Council as the Superintendent and asks for your "
    "quotation as a declared subconsultant. Section E clause 19 of Council's RFQ requires subconsultants to "
    "be declared at tender, and clause 20 of the contract requires Council's approval to subcontract."
)

INSURANCE = [
    "**Public liability:** $20 million for any one occurrence (contract Annexure item 30.2), with Council noted as an interested party.",
    "**Professional indemnity:** $10 million for any one claim with at least one automatic reinstatement (contract Annexure item 30.4). If your cover is lower, state the amount; JCE will confirm whether it is acceptable.",
    "**Workers compensation:** a current WorkCover Queensland policy for all workers you engage on this project. Council's Stage 1 threshold asks for a WorkCover Queensland certificate.",
    "Certificates of currency for each policy, with expiry dates, attached to your quotation, and the policies kept current to 30/06/2030.",
]
INDEPENDENCE = [
    "Declare any current or past relationship (employment, contract, joint venture or financial interest) with AECOM, with Council staff on this project, or with any party likely to tender for Contracts 1 to 3.",
    "If engaged, you must stay independent of the construction contractors and their subcontractors for the life of the engagement, and must not tender for or carry out work for them on this project.",
    "Notify JCE in writing straight away if a conflict arises later. JCE will declare your position in Council's Schedule of Conflict of Interest.",
]
CONFIDENTIALITY = [
    "This request and Council's documents are confidential. Use them only to prepare your quotation and do not pass them to any other party.",
    "Do not contact Council about TCW00728. Council directs all enquiries through VendorPanel. Send your questions to JCE and JCE will raise them with Council if needed (Council's question cut-off is 5:00 pm AEST Friday 23/10/2026).",
    "JCE will keep your prices confidential and will use your company profile, CVs and certificates only in its quotation to Council. Council has advised that quotations may be disclosed under the Right to Information Act 2009 (Qld) and the Local Government Regulation 2012 (Qld).",
]
CONDITIONS = [
    "Any engagement depends on JCE being awarded the contract by Council and on Council approving your firm and your nominated people. JCE does not pay for the preparation of quotations.",
    "The subcontract will be on terms back to back with Council's amended AS 4122-2010, including payment within 30 business days after a valid payment claim (Annexure item 8), monthly claims, and Council's limit of liability.",
    "Hours are Council's estimates and are not guaranteed. Construction contract dates are Council's indicative program and may move. Additional hours caused by later award or longer construction periods will be paid at your quoted rates if Council pays JCE for them.",
    "Nominated personnel may be changed only with JCE's and Council's approval (contract clause 18.1), with a CV showing equivalent experience.",
]


def common_tail(D, info_extra, returnables, pricing_fn, discipline_label):
    D.h("Insurances")
    D.p("Please hold, and confirm in your quotation, insurances matching Council's contract:", keep=True)
    D.bullets(INSURANCE)
    D.h("Work health and safety and systems")
    D.bullets([
        "Your own WHS management system and site-specific SWMS.",
        "Registration and validation in Council's Contractor Management System and completion of the TCC online safety induction before starting, plus each construction Contractor's site induction.",
        "Compliance with each Contractor's site rules. Neither JCE nor your firm is the Principal Contractor.",
    ])
    D.h("Independence and conflict of interest")
    D.bullets(INDEPENDENCE)
    D.h("Confidentiality")
    D.bullets(CONFIDENTIALITY)
    D.h("Information provided")
    D.p("Sent with this request: Council's Request for Quotation (TCW00728), the amended AS 4122-2010 "
        "consultancy contract and Council's Price Schedule V3. " + info_extra + " Council has not yet issued "
        "Appendix E.1 (design drawings), Appendix E.2 (AS 4000 administration table) or the construction "
        "contracts. JCE will ask Council for them and will pass on any addenda.")
    D.h("What to include in your quotation")
    D.bullets(returnables)
    pricing_fn(D)
    D.h("Conditions")
    D.bullets(CONDITIONS)
    D.h("Submission")
    D.p(f"Please email your quotation, CVs and certificates to jayden@jceinfrastructure.com.au by "
        f"**{QUOTE_DUE}**. JCE needs that week to name your people in Council's schedules before the head "
        f"tender closes at {CLOSE}. Questions to Jayden Callen on 0477 207 027.", keep=True)
    D.signature()


# ---------------------------------------------------------------- Construction verification
def build_cv(path, to_line):
    D = Doc()
    D.letterhead(
        "Superintendent's Representative and Inspectors (Townsville)",
        "Vickers Road PS C6B Diversion Project: Superintendent and Construction Verification Services",
        "JCE-RFQ-VP528755-CV (Final)", to_line)
    D.h("Background")
    D.p(BACKGROUND, keep=True)
    D.table(["Contract", "Works", "Construction period"], CONTRACT_ROWS, [2.4, 10.4, 4.2])
    D.p(BACKGROUND_2)
    D.p(JCE_INTENT)

    D.h("Roles and site presence")
    D.p("Council's scope (RFQ Section E clause 5) is split between JCE and your firm as follows:", keep=True)
    D.table(["Role", "Provided by", "Site presence"], [
        ["Superintendent", "JCE (Jayden Callen). Holds the Superintendent's role under AS 4000 for all three contracts and signs every notice, direction, certificate and determination.",
         "Pre-start, monthly progress meeting, key hold points, shutdowns and tie-ins, PC inspection for each contract in person; other meetings by Microsoft Teams. More time on site in critical phases."],
        ["Superintendent's Representative", "Your firm. Townsville-based, acting under a written delegation from the Superintendent.",
         "Chairs the weekly site meeting and attends site as needed: about 7 hours a week during construction, plus establishment, PC, DLP and finalisation."],
        ["Inspector(s)", "Your firm. Townsville-based.",
         "On site 5 days a week while construction is under way, on average 2.5 hours a day including travel (Council's requirement). A second Inspector covers the second site when contracts overlap."],
        ["M&E verification (Contract 2)", "A separate specialist engaged by JCE.", "Not part of this request."],
    ], [3.6, 6.4, 7.0])

    D.h("Scope of services")
    D.sub("Superintendent's Representative")
    D.bullets([
        "Chair weekly site meetings and issue minutes; act as first point of contact on site.",
        "Manage and coordinate the Inspectors and check their reports before issue.",
        "Prepare draft assessments of payment claims, measured quantities, variations and EOT claims for the Superintendent, with the supporting evidence.",
        "Review the Contractor's ITPs, EMP, WHS, QA and traffic management plans, program and work methods (including shutdown, bypass and trenchless methods) and give comments to the Superintendent.",
        "Attend pre-start, dilapidation, PC and defects inspections; maintain the defects list during the DLP.",
        "Contribute to the Surveillance Plan, the monthly progress reports and the Completion Report for each contract.",
    ])
    D.sub("Inspectors")
    D.bullets([
        "Daily technical inspections while work is under way, with a site inspection report for each day (JCE's template).",
        "Witness hold and witness points and sign off ITPs for all work under the contract.",
        "Check measurements and quantities; identify misread or overlooked items in the drawings and specifications.",
        "Witness PE butt welding and electrofusion (including welder qualifications and weld logs), bedding, backfill and compaction testing, thrust blocks, valves, envelope and carrier installation, and hydrostatic pressure testing.",
        "Raise and track NCRs; monitor the Contractor's compliance with its QA, safety, environmental and traffic management plans.",
        "Keep a daily diary and photographs referenced to chainage, and issue diaries, reports and photographs weekly.",
    ])
    D.sub("Records and systems")
    D.bullets([
        "Keep all records in JCE's shared project register in the formats JCE sets (report templates, registers, photo naming). Records are Council's property.",
        "Work to JCE's Surveillance Plan for each contract.",
    ])

    D.h("Indicative hours and program")
    D.p("JCE's split of Council's fixed hours (Price Schedule V3). Council's totals are 3,135 Superintendent "
        "and Representative hours and 2,717 Inspector hours; about 150 Inspector hours on Contract 2 go to "
        "the M&E specialist.", keep=True)
    D.table(["Contract", "Service period", "Superintendent's Representative (h)", "Inspector (h)"], [
        ["Contract 1", "12/2026 to 01/2028, then DLP", "410", "793"],
        ["Contract 2", "07/2027 to 11/2028, then DLP", "500", "812"],
        ["Contract 3", "03/2028 to 07/2029, then DLP", "500", "962"],
        ["Total", "To 30/06/2030", "1,410", "2,567"],
    ], [2.6, 5.4, 5.0, 4.0], bold_last=True)
    D.p("Council's Inspector hours for each contract include 8 hours for establishment, 20 hours for PC and "
        "DLP and 8 hours for finalisation; the rest is construction phase. Overlaps: Contracts 1 and 2 run "
        "together from July 2027 to January 2028, and Contracts 2 and 3 from March to November 2028.")

    D.h("Deliverables")
    D.table(["Deliverable", "Format", "Timing"], [
        ["Site inspection reports and Inspector daily diaries", "PDF", "Weekly to JCE"],
        ["Photographs referenced to chainage", "JPEG", "Weekly"],
        ["Draft assessments of claims, quantities, variations and EOTs", "Excel and PDF", "Within the times JCE sets to meet the contract periods"],
        ["Site meeting agendas and minutes", "Word or PDF", "Before and within 2 working days after each meeting"],
        ["NCR, ITP and defects registers", "Excel", "Kept current; issued monthly"],
        ["Input to Surveillance Plan, monthly reports and Completion Report", "Word", "As JCE programs them"],
    ], [8.0, 3.0, 6.0])

    D.h("Qualifications and RPEQ")
    D.bullets([
        "The services are carried out in Queensland. Under the Professional Engineers Act 2002 (Qld), professional engineering services must be carried out by a registered professional engineer of Queensland (RPEQ) or under the direct supervision of an RPEQ. Council's RFQ does not name a registration, but JCE's Superintendent is not an RPEQ, so this requirement must be met within your team.",
        "The Superintendent's Representative should be an RPEQ (civil). If not, name the RPEQ in your firm who will directly supervise the Representative's and Inspectors' engineering work, with their registration number.",
        "Inspectors: relevant civil construction inspection experience on pressure mains and pump stations; PE welding inspection knowledge; a current general construction induction (White Card) and first aid.",
        "Council scores the nominated individuals on: hands-on superintendency on major pipeline or linear infrastructure; contract administration of sewage pump station M&E and odour control; administering AS 4000-series contracts for local or State government or a major water authority; EOT and variation assessment and dispute avoidance; and backfill depth. Please show this evidence in the CVs.",
    ])

    def pricing(D):
        D.h("Pricing schedule")
        D.p("Prices in AUD, excluding GST. Rates must be all-inclusive, as Council requires: vehicle fitted "
            "with beacon, reversing beeper and UHF radio, fuel and running costs, travel to and from site, "
            "phone, laptop and software, PPE, tools and surveillance equipment. Council pays no separate "
            "disbursements.", keep=True)
        D.table(["Item", "Description", "Hours", "Rate (ex GST)", "Amount (ex GST)"], [
            ["A1", "Hourly rate, Superintendent's Representative", "", "$______ /h", ""],
            ["A2", "Hourly rate, Inspector", "", "$______ /h", ""],
            ["A3", "Hourly rate, backup Representative or backup Inspector (if different)", "", "$______ /h", ""],
            ["A4", "Hourly rate, supervising RPEQ (if separate from the Representative)", "", "$______ /h", ""],
            ["B1", "Contract 1: Superintendent's Representative (lump sum, not to exceed)", "410", "A1", "$__________"],
            ["B2", "Contract 1: Inspector (lump sum, not to exceed)", "793", "A2", "$__________"],
            ["B3", "Contract 2: Superintendent's Representative (lump sum, not to exceed)", "500", "A1", "$__________"],
            ["B4", "Contract 2: Inspector (lump sum, not to exceed)", "812", "A2", "$__________"],
            ["B5", "Contract 3: Superintendent's Representative (lump sum, not to exceed)", "500", "A1", "$__________"],
            ["B6", "Contract 3: Inspector (lump sum, not to exceed)", "962", "A2", "$__________"],
            ["B7", "Any other cost not covered by the rates (specify)", "", "", "$__________"],
            ["", "Total, items B1 to B7 (excluding GST)", "3,977", "", "$__________"],
        ], [1.2, 8.0, 1.6, 2.8, 3.4], bold_last=True)
        D.bullets([
            "**Exclusions:** list every exclusion and assumption. Anything not listed is taken as included.",
            f"**Validity:** prices to remain valid for at least 90 days from {QUOTE_DUE.split(' ', 3)[-1]} (to {VALID_TO}). Council's quotation validity is 150 days from 30/10/2026; please say whether you will extend if Council's award is later than {VALID_TO}.",
            "**Rates over time:** state whether rates hold to 30/06/2030 or give the annual escalation basis (for example a fixed percentage each 1 July).",
            "**Hours:** state any minimum call-out or minimum daily charge. Council pays for hours worked only.",
        ])

    common_tail(D, "JCE's draft site inspection report format is available on request.", [
        "Completed pricing schedule (section 13), exclusions and validity.",
        "CVs (2 pages each) for the proposed Superintendent's Representative, at least two Inspectors and named backups, with RPEQ numbers where held and two referees each.",
        "A short statement of how you will cover the Inspector's 5-day site presence, the overlap periods, leave and illness.",
        "Company profile; Townsville office address and post code; number of Townsville-resident FTE (for Council's Schedule of Local Industry Participation).",
        "Similar projects in the last 5 years (client, contact, scope, value, period), and any work for Townsville City Council in the last 3 years with a Council referee.",
        "Certificates of currency (public liability, professional indemnity, WorkCover Queensland).",
        "Conflict of interest declaration (section 9) and confirmation of WHS system and SWMS.",
        "Any social benefit or Indigenous engagement points (for example Indigenous ownership, apprentices or trainees, local community contribution).",
    ], pricing, "CV")
    D.footer("JCE Infrastructure  |  RFQ JCE-RFQ-VP528755-CV  |  TCW00728 Vickers Rd PS C6B  |  Issued 12/10/2026")
    D.save(path)


# ---------------------------------------------------------------- M&E verification
def build_me(path, to_line):
    D = Doc()
    D.letterhead(
        "Mechanical and Electrical Verification, PS C6B (Contract 2)",
        "Vickers Road PS C6B Diversion Project: Superintendent and Construction Verification Services",
        "JCE-RFQ-VP528755-ME (Final)", to_line)
    D.h("Background")
    D.p(BACKGROUND, keep=True)
    D.table(["Contract", "Works", "Construction period"], CONTRACT_ROWS, [2.4, 10.4, 4.2])
    D.p(BACKGROUND_2)
    D.p(JCE_INTENT)

    D.h("Roles and site presence")
    D.table(["Role", "Provided by", "Site presence"], [
        ["Superintendent", "JCE (Jayden Callen). Signs every notice, direction, certificate and determination, and releases hold points.",
         "Key hold points, shutdowns and PC inspection in person."],
        ["Superintendent's Representative and Inspectors", "A Townsville partner engaged by JCE.", "Daily inspections while work is under way."],
        ["M&E verification (Contract 2 only)", "Your firm. Reviews and witnesses the mechanical, electrical, controls and odour control work and recommends hold point releases to the Superintendent.",
         "Called off by JCE to match the Contractor's submittals and test dates: about 8 to 12 attendances at PS C6B (North Vickers Rd, Condon) between July 2027 and November 2028, plus factory tests if Council requires them. Document review can be done remotely."],
    ], [3.6, 6.6, 6.8])

    D.h("Scope of services")
    D.bullets([
        "**Submittal review:** review the Contractor's mechanical, electrical and controls submittals for PS C6B (pumps, valves and pipework, switchboard and MCC drawings, SCADA and telemetry functional description, odour control unit design to AECOM's performance specification) against the contract specification and TCC standard specifications, and give written comments to the Superintendent.",
        "**Factory acceptance tests:** witness switchboard and pump FATs where Council requires them.",
        "**Installation inspections:** inspect the mechanical and electrical installation at the hold and witness points in the Contractor's ITPs.",
        "**Site acceptance and commissioning:** witness site acceptance tests and SCADA and telemetry commissioning with Council's operations staff.",
        "**Odour control unit:** witness performance testing against AECOM's performance specification, and review commissioning records and O&M manuals before Practical Completion.",
        "**Reporting:** a short report for each attendance, recommended hold point releases, and input to the PC defects list and the Contract 2 Completion Report.",
    ])

    D.h("Indicative hours and program")
    D.p("About 150 hours across Contract 2 (July 2027 to November 2028), within Council's Inspector "
        "allowance. JCE will call off attendances as needed; hours are not guaranteed.", keep=True)
    D.table(["Activity", "Indicative hours"], [
        ["Submittal review (mechanical, electrical, SCADA, odour control unit)", "40"],
        ["Factory acceptance test witnessing (if Council requires FAT)", "30"],
        ["Installation inspections at hold and witness points", "30"],
        ["Site acceptance tests and SCADA and telemetry commissioning", "30"],
        ["Odour control unit performance testing and O&M review", "20"],
        ["Total", "150"],
    ], [13.0, 4.0], bold_last=True)

    D.h("Deliverables")
    D.table(["Deliverable", "Format", "Timing"], [
        ["Submittal review comments", "Excel or PDF register", "Within 5 working days of receipt"],
        ["Inspection and test attendance reports with recommended hold point releases", "PDF", "Within 2 working days of the attendance"],
        ["Input to the PC defects list and Contract 2 Completion Report", "Word or Excel", "At PC and at finalisation"],
    ], [8.0, 4.0, 5.0])

    D.h("Qualifications and RPEQ")
    D.bullets([
        "The services are carried out in Queensland. Under the Professional Engineers Act 2002 (Qld), professional engineering services must be carried out by an RPEQ or under the direct supervision of an RPEQ. Council's RFQ does not name a registration, but JCE's Superintendent is not an RPEQ.",
        "The engineer who signs submittal reviews and test reports must be an RPEQ in the electrical or mechanical area, or work under the direct supervision of a named RPEQ in your firm. Please give registration numbers.",
        "Experience wanted: sewage pump station mechanical and electrical works, switchboards and MCCs, SCADA and telemetry, and odour control commissioning for local government or water authorities.",
    ])

    def pricing(D):
        D.h("Pricing schedule")
        D.p("Prices in AUD, excluding GST. Rates must be all-inclusive (vehicle, fuel, phone, laptop and "
            "software, PPE, tools, and travel within Townsville). Council pays no separate disbursements, so "
            "price any travel from outside Townsville in item B6 or B7.", keep=True)
        D.table(["Item", "Description", "Hours", "Rate (ex GST)", "Amount (ex GST)"], [
            ["A1", "Hourly rate, senior electrical or controls engineer (RPEQ)", "", "$______ /h", ""],
            ["A2", "Hourly rate, mechanical engineer (RPEQ)", "", "$______ /h", ""],
            ["A3", "Hourly rate, backup engineer (if different)", "", "$______ /h", ""],
            ["B1", "Submittal review (lump sum, not to exceed)", "40", "", "$__________"],
            ["B2", "FAT witnessing, if required (lump sum, excluding travel)", "30", "", "$__________"],
            ["B3", "Installation inspections (lump sum, not to exceed)", "30", "", "$__________"],
            ["B4", "SAT and SCADA and telemetry commissioning (lump sum, not to exceed)", "30", "", "$__________"],
            ["B5", "Odour control unit performance testing and O&M review (lump sum, not to exceed)", "20", "", "$__________"],
            ["B6", "Travel and expenses for each FAT day outside Townsville, if required", "", "$______ /day", "$__________"],
            ["B7", "Travel to PS C6B from your base, if outside Townsville (per attendance)", "", "$______ /trip", "$__________"],
            ["", "Total, items B1 to B5 (excluding GST)", "150", "", "$__________"],
        ], [1.2, 8.0, 1.6, 2.8, 3.4], bold_last=True)
        D.bullets([
            "**Exclusions:** list every exclusion and assumption. Anything not listed is taken as included.",
            f"**Validity:** prices to remain valid for at least 90 days from 23/10/2026 (to {VALID_TO}). Council's quotation validity is 150 days from 30/10/2026; please say whether you will extend if Council's award is later than {VALID_TO}.",
            "**Rates over time:** the work falls between July 2027 and November 2028. State whether rates hold for that period or give the escalation basis.",
            "**Hours:** state any minimum call-out or minimum daily charge.",
        ])

    common_tail(D, "", [
        "Completed pricing schedule (section 13), exclusions and validity.",
        "CVs (2 pages each) for the proposed engineer or engineers and a backup, with RPEQ numbers and two referees each.",
        "Company profile, office address and post code, and any work for Townsville City Council in the last 3 years.",
        "Similar projects in the last 5 years (client, contact, scope, value, period).",
        "Certificates of currency (public liability, professional indemnity, WorkCover Queensland or, for staff based outside Queensland, the relevant workers compensation policy).",
        "Conflict of interest declaration (section 9) and confirmation of WHS system and SWMS.",
    ], pricing, "ME")
    D.footer("JCE Infrastructure  |  RFQ JCE-RFQ-VP528755-ME  |  TCW00728 Vickers Rd PS C6B  |  Issued 12/10/2026")
    D.save(path)


if __name__ == "__main__":
    generic = "Invited subconsultant (named in the covering email)"
    jobs = [
        ("Construction_Verification", build_cv, ["PDM Managers", "Gnosis Engineering", "Steve McKenzie Consultant Engineer Pty Ltd (SMCE)"]),
        ("ME_Verification", build_me, ["Orana Engineering", "Clarence Consultants (attention Ken Heard)"]),
    ]
    for disc, fn, firms in jobs:
        d = os.path.join(BASE, disc)
        fn(os.path.join(d, f"JCE_RFQ_{disc}_Vickers-Rd-PS-C6B_Final.docx"), generic)
        for firm in firms:
            slug = re.sub(r"[^A-Za-z0-9]+", "-", firm.split(" (")[0]).strip("-")
            fn(os.path.join(d, "per_firm", f"JCE_RFQ_{disc}_Vickers-Rd-PS-C6B_Final_{slug}.docx"), firm)
    print("done")
