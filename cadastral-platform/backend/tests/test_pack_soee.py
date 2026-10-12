"""Draft SoEE: deterministic mode, TO BE COMPLETED for unresolved items, pluggable LLM with guard rails, DOCX/PDF output."""
import json

import pytest

from app.agents.sa11_pathway import assess_pathway, make_input
from app.pack.compliance import ProposalFacts, build_compliance_table
from app.pack.samples import ROOT, pathway_for_run
from app.pack.soee import TBC, SoeeInputs, build_soee, write_docx, write_pdf


def _report(run):
    p = ROOT / run / "constraints_report.json"
    if not p.exists():
        pytest.skip("test-run outputs not present")
    return json.loads(p.read_text())


def _inputs(**kw):
    rep = _report("lot2_dwelling")
    facts = ProposalFacts(dev_type="dwelling_house", zone="R2", lga="Gunnedah", lot_area_m2=955.8, constraints=rep, bushfire_prone=True, flood_affected=False,
                          setbacks_m={"front": 7.0, "side": 0.5}, building_height_m=6.5)
    pr = assess_pathway(make_input("dwelling_house", rep, lga="Gunnedah", zone="R2", lot_area_m2=955.8))
    return SoeeInputs(facts=facts, pathway_result=pr, compliance=build_compliance_table(facts, "da"), constraints=rep, lotplan="2//DP227508", **kw)


class FakeLlm:
    def __init__(self, mode="echo"):
        self.mode, self.calls = mode, []

    def complete(self, system, user):
        self.calls.append(user)
        if self.mode == "echo":
            return " ".join(l.split(") ", 1)[-1] for l in user.splitlines())
        if self.mode == "invent":
            return user.replace("2//DP227508", "2//DP227508") + " The site is 1234 square metres."
        if self.mode == "drop_tbc":
            return "Everything is complete."
        raise RuntimeError("down")


def test_deterministic_draft_is_complete_and_plain():
    d = build_soee(_inputs())
    assert d.mode == "deterministic"
    assert [s.number for s in d.sections] == ["1", "2", "3", "4", "5", "6"]
    titles = [s.title for s in d.sections]
    assert titles == ["Site and context", "Proposal description", "Zoning and permissibility", "Relevant planning controls and compliance summary", "Environmental impacts", "Conclusion"]
    text = d.to_text()
    assert "955.8 square metres" in text and "R2" in text and "6.5 m" in text
    assert TBC in text and len(d.tbc_items()) > 10
    # unresolved items are TO BE COMPLETED, never invented
    assert any("permissibility" in s.text for s in d.tbc_items())
    assert any("Site address" in s.text and s.tbc for s in d.sections[0].statements)
    for sec in d.sections:
        for s in sec.statements:
            assert s.source, s.id


def test_failed_control_is_stated_with_variation_to_be_completed():
    d = build_soee(_inputs())
    s4 = " ".join(s.text for s in d.sections[3].statements)
    assert "Control not met: Setback: side" in s4 and "Variation justification: TO BE COMPLETED" in s4
    assert d.sections[3].table and d.sections[3].table[0][0] == "Control"


def test_every_statement_is_tied_to_a_field_and_no_approval_language():
    t = build_soee(_inputs()).to_text().lower()
    assert "is approved" not in t and "complies with" not in t.replace("compliance", "")


def test_provided_facts_replace_to_be_completed():
    d = build_soee(_inputs(site_address="1 Test St, Gunnedah NSW 2380", proposal_description="Single storey dwelling", lep_permissibility="permitted_with_consent", escp_reference="ESCP-1"))
    t = d.to_text()
    assert "1 Test St, Gunnedah NSW 2380" in t and "Single storey dwelling" in t and "permitted with consent" in t and "ESCP-1" in t
    assert not any("Site address" in s.text for s in d.tbc_items())


def test_llm_prose_accepted_when_it_adds_nothing():
    llm = FakeLlm("echo")
    d = build_soee(_inputs(), llm)
    assert d.mode == "llm" and len(llm.calls) == 6 and not d.warnings
    assert TBC in d.to_text()


def test_llm_prose_with_invented_number_or_dropped_tbc_is_rejected():
    for mode in ("invent", "drop_tbc", "error"):
        d = build_soee(_inputs(), FakeLlm(mode))
        assert d.mode == "deterministic" and d.warnings, mode
        assert "1234" not in d.to_text()


def test_docx_and_pdf_open_with_watermark_and_disclaimer(tmp_path):
    from docx import Document
    d = build_soee(_inputs())
    dx, pf = write_docx(d, tmp_path / "s.docx"), write_pdf(d, tmp_path / "s.pdf")
    doc = Document(str(dx))
    body = "\n".join(p.text for p in doc.paragraphs)
    assert "TO BE COMPLETED" in body and "drafting aid" in body and "suitably qualified person" in body
    assert "DRAFT" in doc.sections[0].header.paragraphs[0].text
    assert "PowerPlusWaterMarkObject" in doc.sections[0].header._element.xml and 'string="DRAFT"' in doc.sections[0].header._element.xml
    assert len(doc.tables) == 1
    import fitz
    pdf = fitz.open(str(pf))
    txt = "".join(p.get_text() for p in pdf)
    assert "TO BE COMPLETED" in txt and "drafting aid" in txt and "DRAFT" in txt and len(pdf) >= 3


def test_real_quia_granny_draft():
    rep = _report("quia_granny")
    pr, r = pathway_for_run("quia_granny")
    facts = ProposalFacts(dev_type="secondary_dwelling", zone="RU1", lga="Gunnedah", lot_area_m2=509264.46, constraints=rep, bushfire_prone=True)
    d = build_soee(SoeeInputs(facts=facts, pathway_result=pr, compliance=build_compliance_table(facts, "da"), constraints=rep, lotplan="3//DP263990"))
    t = d.to_text()
    assert "secondary dwelling (granny flat)" in t and "509264" in t and TBC in t and "Vegetation Category 3" in t
