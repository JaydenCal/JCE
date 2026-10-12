"""Draft Statement of Environmental Effects (SoEE): drafting aid only.

* Facts come only from structured inputs (ProposalFacts, SoeeInputs, SA9 constraints, compliance table, SA11 pathway result).
* Every statement records its `source` (the input field / table row it came from).  Unknown items are written as
  'TO BE COMPLETED: <what is needed>' and never invented.
* Deterministic template mode needs no LLM and produces a complete, plain draft.  An optional LlmClient may rewrite the prose of a
  section; its output is accepted only if it introduces no number absent from the section's facts and keeps every
  'TO BE COMPLETED' line, otherwise the deterministic text is used.
* Structure: Gunnedah Shire Council's SoEE form headings (updated 15 Aug 2023, obtained) and the content a DA must include per the
  NSW Planning Portal 'Local development' page (obtained).  EP&A Regulation 2021 content requirements could not be obtained.
Output: DOCX (python-docx) and PDF (reportlab), watermarked DRAFT, with a front disclaimer.
"""
from __future__ import annotations

import re
from datetime import date
from pathlib import Path
from typing import Any, Literal, Protocol

from pydantic import BaseModel, ConfigDict, Field

from app.agents.requirements.models import Source
from app.agents.sa11_pathway.models import PathwayResult, Status

from .compliance import ComplianceTable, ProposalFacts

TBC = "TO BE COMPLETED"
DISCLAIMER = ("This Statement of Environmental Effects has been prepared as a drafting aid. It must be reviewed, completed and adopted by "
              "a suitably qualified person before it is relied on or lodged. It is not a planning assessment, an approval or advice. "
              "Items marked 'TO BE COMPLETED' have not been determined from the data supplied.")
STRUCTURE_NOTE = ("Section headings follow Gunnedah Shire Council's Statement of Environmental Effects form (last updated 15 August 2023) and the "
                  "NSW Planning Portal 'Local development' requirement for an environmental assessment such as a statement of environmental effects. "
                  "The content requirements in the Environmental Planning and Assessment Regulation 2021 could not be obtained (legislation.nsw.gov.au is blocked): "
                  "TO BE COMPLETED: check the draft against the Regulation.")
DEV_LABEL = {"dwelling_house": "a dwelling house", "secondary_dwelling": "a secondary dwelling (granny flat)", "shed": "a shed / outbuilding"}

SRC_FORM = Source(document_title="Gunnedah Shire Council Statement of Environmental Effects (single dwelling houses, residential ancillary & outbuilding developments)",
                  authority="Gunnedah Shire Council", version="Last updated 15 August 2023", page=1, quote="Statement of Environmental Effects is to be submitted with all development applications other than",
                  sha256="", source_url=None)


class LlmClient(Protocol):
    def complete(self, system: str, user: str) -> str: ...


class M(BaseModel):
    model_config = ConfigDict(extra="forbid", arbitrary_types_allowed=True)


class SoeeInputs(M):
    facts: ProposalFacts
    pathway_result: PathwayResult | None = None
    compliance: ComplianceTable | None = None
    constraints: dict | None = None
    project_title: str | None = None
    site_address: str | None = None
    lotplan: str | None = None
    council_name: str | None = None
    applicant_name: str | None = None
    owner_name: str | None = None
    prepared_by: str | None = None
    proposal_description: str | None = None
    estimated_cost_aud: float | None = None
    bedrooms: int | None = None
    external_materials: str | None = None
    existing_use: str | None = None
    previous_uses: str | None = None
    adjoining_uses: str | None = None
    topography: str | None = None
    vegetation: str | None = None
    contamination_history: str | None = None
    electricity_source: str | None = None
    water_source: str | None = None
    sewer_or_effluent: str | None = None
    stormwater_disposal: str | None = None
    access_road_name: str | None = None
    overshadowing_privacy_comment: str | None = None
    waste_management: str | None = None
    escp_reference: str | None = None            # ESCP (SA10) document reference if present
    heritage_comment: str | None = None
    lep_permissibility: Literal["permitted_with_consent", "permitted_without_consent", "prohibited"] | None = None
    today: date = Field(default_factory=date.today)


class Statement(M):
    id: str
    text: str
    source: str                        # input field / table row this statement is derived from
    tbc: bool = False
    citation: Source | None = None


class Section(M):
    number: str
    title: str
    statements: list[Statement]
    table: list[list[str]] | None = None        # optional compliance table (header + rows)
    prose: str | None = None                    # LLM prose (only when validated)


class SoeeDocument(M):
    title: str
    mode: Literal["deterministic", "llm"]
    sections: list[Section]
    warnings: list[str] = Field(default_factory=list)
    disclaimer: str = DISCLAIMER
    structure_note: str = STRUCTURE_NOTE
    generated: str

    def tbc_items(self) -> list[Statement]:
        return [s for sec in self.sections for s in sec.statements if s.tbc]

    def to_text(self) -> str:
        out = [self.title, "", "DRAFT", self.disclaimer, ""]
        for sec in self.sections:
            out.append(f"{sec.number}. {sec.title}")
            out.append(sec.prose if sec.prose else "\n".join(s.text for s in sec.statements))
            if sec.table:
                out += [" | ".join(r) for r in sec.table]
            out.append("")
        return "\n".join(out)


class _B:
    def __init__(self, inp: SoeeInputs) -> None:
        self.i, self.n, self.items = inp, 0, []

    def fact(self, text: str, source: str, citation: Source | None = None) -> None:
        self.n += 1; self.items.append(Statement(id=f"s{self.n}", text=text, source=source, citation=citation))

    def val(self, label: str, value: Any, source: str, *, unit: str = "", need: str | None = None, fmt: str = "{label}: {v}.") -> None:
        if value is None or value == "":
            self.n += 1
            self.items.append(Statement(id=f"s{self.n}", text=f"{label}: {TBC} ({need or 'to be supplied by the author'}).", source=source, tbc=True))
        else:
            v = f"{value:g}" if isinstance(value, float) else str(value)
            self.fact(fmt.format(label=label, v=f"{v} {unit}".strip()), source)

    def tbc(self, text: str, source: str) -> None:
        self.n += 1; self.items.append(Statement(id=f"s{self.n}", text=f"{TBC}: {text}", source=source, tbc=True))


def _layer(c: dict | None, key: str) -> dict | None:
    return next((l for l in (c or {}).get("layers", []) if l["key"] == key), None)


def _layer_stmt(b: _B, label: str, key: str, c: dict | None, need: str) -> None:
    l = _layer(c, key)
    src = f"constraints.{key}"
    if l is None or l["status"] == "not_checked":
        b.tbc(f"{label}: the constraints layer was not checked ({need}).", src)
    elif l["status"] == "affected":
        attrs = "; ".join(", ".join(f"{k}={v}" for k, v in a.items() if v not in (None, "")) for a in l["attributes"][:3])
        b.fact(f"{label}: the NSW mapping layer reports the lot as affected ({attrs}). Data currency: {l.get('data_currency', 'not stated')[:120]}.", src)
    else:
        b.fact(f"{label}: the NSW mapping layer reports the lot as not affected. This is a map layer result, not a planning certificate.", src)


def build_soee(inp: SoeeInputs, llm: LlmClient | None = None) -> SoeeDocument:
    f, c = inp.facts, inp.constraints
    secs: list[Section] = []

    def sec(num: str, title: str, fn) -> None:
        b = _B(inp); fn(b); secs.append(Section(number=num, title=title, statements=b.items))

    # 1 Site and context
    def s1(b: _B) -> None:
        b.val("Site address", inp.site_address, "inputs.site_address", need="street address")
        b.val("Lot and plan", inp.lotplan, "inputs.lotplan")
        b.val("Local government area", f.lga or inp.council_name, "facts.lga")
        b.val("Zone", f.zone, "facts.zone", need="LEP zone from the planning certificate", fmt="{label}: {v} (from the lot data supplied).")
        b.val("Lot area", f.lot_area_m2, "facts.lot_area_m2", unit="square metres", fmt="{label}: {v} (calculated from the cadastre, not a survey).")
        b.val("Lot width at the building line", f.lot_width_m, "facts.lot_width_m", unit="m")
        b.val("Topography (slope, vegetation, orientation, streetscape)", inp.topography, "inputs.topography")
        b.val("Existing vegetation", inp.vegetation, "inputs.vegetation")
        b.val("Present use of the site", inp.existing_use, "inputs.existing_use")
        b.val("Previous uses of the site", inp.previous_uses, "inputs.previous_uses")
        b.val("Use of adjoining land", inp.adjoining_uses, "inputs.adjoining_uses")
        b.val("Potentially contaminating land uses (current or previous)", inp.contamination_history, "inputs.contamination_history", need="answer the contamination questions on the Council form")
        _layer_stmt(b, "Contaminated land", "nsw_contamination", c, "obtain the planning certificate")
    sec("1", "Site and context", s1)

    # 2 Proposal
    def s2(b: _B) -> None:
        b.fact(f"The proposal is {DEV_LABEL[f.dev_type]}.", "facts.dev_type")
        b.val("Description of the development", inp.proposal_description, "inputs.proposal_description", need="description of the development")
        b.val("Estimated cost of the development", inp.estimated_cost_aud, "inputs.estimated_cost_aud", unit="AUD", need="required in the DA (NSW Planning Portal)", fmt="{label}: {v}.")
        b.val("Storeys", f.storeys, "facts.storeys")
        b.val("Building height", f.building_height_m, "facts.building_height_m", unit="m")
        b.val("Gross floor area" if f.dev_type != "shed" else "Floor area", f.gfa_m2 if f.dev_type != "shed" else f.floor_area_m2,
              "facts.gfa_m2" if f.dev_type != "shed" else "facts.floor_area_m2", unit="square metres")
        for k in ("front", "side", "rear"):
            b.val(f"Setback to the {k} boundary", f.setbacks_m.get(k), f"facts.setbacks_m.{k}", unit="m", need="from the site plan")
        b.val("Car parking spaces", f.parking_spaces, "facts.parking_spaces", need="from the site plan")
        b.val("External materials and finishes", inp.external_materials, "inputs.external_materials", need="Council requires samples of colours and materials")
    sec("2", "Proposal description", s2)

    # 3 Zoning and permissibility
    def s3(b: _B) -> None:
        b.val("Zone", f.zone, "facts.zone")
        perm = inp.lep_permissibility or (inp.pathway_result and None)
        if inp.lep_permissibility:
            b.fact(f"Permissibility under the LEP land use table: {inp.lep_permissibility.replace('_', ' ')} (entered by the author; LEP text not checked by this tool).", "inputs.lep_permissibility")
        else:
            b.tbc("permissibility of the development in the zone under the Local Environmental Plan land use table (the LEP was not obtained)", "inputs.lep_permissibility")
        pr = inp.pathway_result
        if pr:
            for card in pr.cards:
                b.fact(f"Approval pathway check, {card.title}: {card.pill.value}. {card.headline}", f"pathway.{card.pathway}")
            b.fact(f"{pr.recommendation} This is decision support, not an approval.", "pathway.recommendation")
        else:
            b.tbc("approval pathway assessment (SA11 result not supplied)", "pathway")
        b.tbc("compatibility with the zone objectives and with adjoining development (LEP zone objectives not obtained)", "lep.zone_objectives")
    sec("3", "Zoning and permissibility", s3)

    # 4 Planning controls and compliance summary
    def s4(b: _B) -> None:
        t = inp.compliance
        if t is None:
            b.tbc("compliance table not supplied", "compliance"); return
        n = t.counts()
        b.fact(f"Compliance summary for the {t.pathway.upper()} pathway: {n['PASS']} controls pass, {n['FAIL']} fail, {n['UNRESOLVED']} are unresolved and {n['NOT_APPLICABLE']} are not applicable.", "compliance.counts")
        for r in t.rows:
            if r.status == Status.FAIL:
                b.fact(f"Control not met: {r.control} (required {r.requirement_text}; proposed {r.proposed_value}). Variation justification: {TBC}.", f"compliance.{r.control}", r.citation)
            elif r.status == Status.UNRESOLVED:
                b.tbc(f"{r.control}: {r.prompt or 'not determined'}", f"compliance.{r.control}")
    sec("4", "Relevant planning controls and compliance summary", s4)
    if inp.compliance:
        secs[-1].table = [["Control", "Requirement and source", "Proposed", "Status"]] + [
            [r.control, r.requirement_text + (f" [{r.citation.document_title}, p.{r.citation.page}: \"{r.citation.quote[:110]}\"]" if r.citation else ""),
             "" if r.proposed_value is None else str(r.proposed_value), r.status.value.replace("_", " ")] for r in inp.compliance.rows]

    # 5 Environmental impacts
    def s5(b: _B) -> None:
        b.fact("Flood", "heading")
        fc = f.flood
        if fc is not None:
            b.fact(f"Flood check result: {fc.result}. Minimum floor level {fc.min_ffl} m AHD against a proposed floor level of {fc.proposed_ffl} m AHD." if fc.min_ffl is not None and fc.proposed_ffl is not None
                   else f"Flood check result: {fc.result}.", "facts.flood")
        _layer_stmt(b, "Flood planning area", "nsw_flood", c, "obtain flood information from the council")
        b.fact("Bushfire", "heading")
        _layer_stmt(b, "Bushfire prone land", "nsw_bushfire", c, "obtain the planning certificate")
        b.val("Bushfire attack level", f.bal, "facts.bal", need="BAL assessment by an accredited practitioner (AS 3959)")
        b.fact("Heritage", "heading")
        _layer_stmt(b, "Heritage item", "nsw_heritage_item", c, "check the LEP heritage schedule")
        _layer_stmt(b, "Heritage conservation area", "nsw_heritage_conservation", c, "check the LEP heritage map")
        if inp.heritage_comment:
            b.fact(inp.heritage_comment, "inputs.heritage_comment")
        b.fact("Amenity (overshadowing, privacy, views)", "heading")
        b.val("Effect on amenity of surrounding residences", inp.overshadowing_privacy_comment, "inputs.overshadowing_privacy_comment", need="overshadowing (9 am and 3 pm, 21 June, two-storey only) and privacy")
        b.fact("Access and parking", "heading")
        b.val("Road the site is accessed from", inp.access_road_name, "inputs.access_road_name")
        b.val("Crossover width", f.crossover_width_m, "facts.crossover_width_m", unit="m")
        b.val("Car parking spaces", f.parking_spaces, "facts.parking_spaces")
        b.fact("Stormwater", "heading")
        b.val("Stormwater disposal", inp.stormwater_disposal or f.stormwater_strategy, "inputs.stormwater_disposal", need="street, on-site retention or detention, per the Engineering Design Minimum Standards")
        b.fact("Erosion and sediment control", "heading")
        if inp.escp_reference:
            b.fact(f"An erosion and sediment control plan is included in the application pack ({inp.escp_reference}).", "inputs.escp_reference")
        else:
            b.tbc("erosion and sediment control plan (Council's Fact Sheet 1 lists it as required for construction works to minimise harm to the environment)", "inputs.escp_reference")
        b.fact("Services and waste", "heading")
        b.val("Electricity supply", inp.electricity_source, "inputs.electricity_source")
        b.val("Water supply (town supply, rainwater tank, bore)", inp.water_source, "inputs.water_source")
        b.val("Effluent disposal (sewer or on-site)", inp.sewer_or_effluent, "inputs.sewer_or_effluent")
        b.val("Waste management", inp.waste_management, "inputs.waste_management")
        for label, key in (("Acid sulfate soils", "nsw_acid_sulfate"), ("Biodiversity", "nsw_biodiversity"), ("Riparian land", "nsw_riparian")):
            _layer_stmt(b, label, key, c, "check the LEP map")
    sec("5", "Environmental impacts", s5)

    # 6 Conclusion
    def s6(b: _B) -> None:
        t = inp.compliance
        if t:
            n = t.counts()
            b.fact(f"Against the controls tabulated in section 4, {n['PASS']} are met, {n['FAIL']} are not met and {n['UNRESOLVED']} remain to be resolved.", "compliance.counts")
        b.tbc("the author's conclusion on the merits of the proposal and its environmental effects", "inputs.prepared_by")
        b.val("Prepared by", inp.prepared_by, "inputs.prepared_by", need="name and qualifications of the suitably qualified person who reviews and adopts this statement")
    sec("6", "Conclusion", s6)

    doc = SoeeDocument(title=inp.project_title or f"Statement of Environmental Effects: {DEV_LABEL[f.dev_type]}", mode="deterministic", sections=secs, generated=inp.today.isoformat())
    if llm is not None:
        _apply_llm(doc, llm)
    return doc


_NUM = re.compile(r"\d+(?:\.\d+)?")
_SYS = ("You rewrite structured facts into plain, neutral planning prose for a draft Statement of Environmental Effects. Use ONLY the supplied statements. "
        "Do not add numbers, claims, opinions or conclusions. Keep every line that contains 'TO BE COMPLETED' verbatim. Do not state that anything is approved or compliant beyond the facts.")


def _apply_llm(doc: SoeeDocument, llm: LlmClient) -> None:
    used = False
    for sec in doc.sections:
        facts = [s for s in sec.statements if s.text and s.source != "heading"]
        user = "\n".join(f"[{s.id}] ({s.source}) {s.text}" for s in facts)
        try:
            prose = llm.complete(_SYS, user)
        except Exception as e:  # noqa: BLE001
            doc.warnings.append(f"section {sec.number}: LLM error {type(e).__name__}; deterministic text used"); continue
        allowed = set(_NUM.findall(user)) | set(_NUM.findall(sec.title)) | {sec.number}
        extra = set(_NUM.findall(prose)) - allowed
        missing = [s.text for s in facts if s.tbc and s.text not in prose]
        if extra or missing:
            doc.warnings.append(f"section {sec.number}: LLM prose rejected ({'new numbers ' + str(sorted(extra)) if extra else 'TO BE COMPLETED lines altered'}); deterministic text used")
            continue
        sec.prose = prose; used = True
    if used:
        doc.mode = "llm"


# ----------------------------------------------------------------------------------------------------------------------
def write_docx(doc: SoeeDocument, path: str | Path) -> Path:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import parse_xml
    from docx.shared import Pt, RGBColor

    d = Document()
    d.styles["Normal"].font.name = "Calibri"; d.styles["Normal"].font.size = Pt(10.5)
    sect = d.sections[0]
    hdr = sect.header.paragraphs[0]
    hdr.text = "DRAFT: Statement of Environmental Effects (drafting aid, not for lodgement until reviewed)"
    hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    wm = ('<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:v="urn:schemas-microsoft-com:vml" '
          'xmlns:o="urn:schemas-microsoft-com:office:office"><w:pict>'
          '<v:shapetype id="_x0000_t136" coordsize="21600,21600" o:spt="136" adj="10800" path="m@7,l@8,m@5,21600l@6,21600e">'
          '<v:path textpathok="t" o:connecttype="custom"/><v:textpath on="t" fitshape="t"/></v:shapetype>'
          '<v:shape id="PowerPlusWaterMarkObject1" o:spid="_x0000_s2049" type="#_x0000_t136" '
          'style="position:absolute;margin-left:0;margin-top:0;width:430pt;height:140pt;rotation:315;z-index:-251658752;'
          'mso-position-horizontal:center;mso-position-horizontal-relative:margin;mso-position-vertical:center;mso-position-vertical-relative:margin" '
          'o:allowincell="f" fillcolor="silver" stroked="f"><v:fill opacity=".5"/><v:textpath style="font-family:&quot;Calibri&quot;;font-size:1pt" string="DRAFT"/></v:shape>'
          '</w:pict></w:r>')
    hdr._p.append(parse_xml(wm))
    t = d.add_heading(doc.title, 0)
    p = d.add_paragraph(); r = p.add_run("DRAFT"); r.bold = True; r.font.size = Pt(20); r.font.color.rgb = RGBColor(0xC0, 0, 0)
    p = d.add_paragraph(); r = p.add_run("Disclaimer. "); r.bold = True; p.add_run(doc.disclaimer)
    d.add_paragraph(f"Generated {doc.generated}; mode: {doc.mode}. " + doc.structure_note).runs[0].italic = True
    d.add_page_break()
    for sec in doc.sections:
        d.add_heading(f"{sec.number}. {sec.title}", 1)
        if sec.prose:
            d.add_paragraph(sec.prose)
        else:
            for s in sec.statements:
                if s.source == "heading":
                    d.add_heading(s.text, 2); continue
                p = d.add_paragraph(style="List Bullet")
                if s.tbc:
                    a, _, rest = s.text.partition(TBC)
                    p.add_run(a)
                    rr = p.add_run(TBC); rr.bold = True; rr.font.color.rgb = RGBColor(0xC0, 0, 0)
                    p.add_run(rest)
                else:
                    p.add_run(s.text)
        if sec.table:
            tb = d.add_table(rows=0, cols=len(sec.table[0])); tb.style = "Table Grid"
            for i, row in enumerate(sec.table):
                cells = tb.add_row().cells
                for cell, txt in zip(cells, row):
                    cell.text = txt
                    if i == 0:
                        for rn in cell.paragraphs[0].runs: rn.bold = True
    if doc.warnings:
        d.add_heading("Drafting notes", 2)
        for w in doc.warnings: d.add_paragraph(w, style="List Bullet")
    path = Path(path); path.parent.mkdir(parents=True, exist_ok=True); d.save(str(path)); return path


def write_pdf(doc: SoeeDocument, path: str | Path) -> Path:
    from xml.sax.saxutils import escape
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    ss = getSampleStyleSheet()
    small = ss["BodyText"].clone("small", fontSize=7.5, leading=9)

    def marks(canvas, _doc):
        canvas.saveState()
        canvas.setFont("Helvetica-Bold", 110); canvas.setFillColorRGB(0.85, 0.85, 0.85)
        canvas.translate(A4[0] / 2, A4[1] / 2); canvas.rotate(45); canvas.drawCentredString(0, -30, "DRAFT")
        canvas.restoreState()
        canvas.setFont("Helvetica", 8); canvas.setFillColorRGB(0.4, 0.4, 0.4)
        canvas.drawString(18 * mm, 10 * mm, "DRAFT drafting aid: to be reviewed and adopted by a suitably qualified person")
        canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, f"Page {_doc.page}")

    path = Path(path); path.parent.mkdir(parents=True, exist_ok=True)
    pdf = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=18 * mm, bottomMargin=18 * mm, title=doc.title)
    E: list = [Paragraph(escape(doc.title), ss["Title"]), Paragraph('<font color="#C00000"><b>DRAFT</b></font>', ss["Heading1"]),
               Paragraph("<b>Disclaimer.</b> " + escape(doc.disclaimer), ss["BodyText"]), Spacer(1, 6),
               Paragraph(f"<i>Generated {doc.generated}; mode: {doc.mode}. {escape(doc.structure_note)}</i>", small), PageBreak()]
    def tbc_markup(text: str) -> str:
        return escape(text).replace(TBC, f'<font color="#C00000"><b>{TBC}</b></font>')
    for sec in doc.sections:
        E.append(Paragraph(f"{sec.number}. {escape(sec.title)}", ss["Heading1"]))
        if sec.prose:
            E.append(Paragraph(tbc_markup(sec.prose), ss["BodyText"]))
        else:
            for s in sec.statements:
                if s.source == "heading":
                    E.append(Paragraph(escape(s.text), ss["Heading3"]))
                else:
                    E.append(Paragraph("&bull; " + tbc_markup(s.text), ss["BodyText"]))
        if sec.table:
            data = [[Paragraph(escape(c), small) for c in row] for row in sec.table]
            tb = Table(data, colWidths=[38 * mm, 78 * mm, 30 * mm, 24 * mm], repeatRows=1)
            tb.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.grey), ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
            E += [Spacer(1, 4), tb]
    if doc.warnings:
        E.append(Paragraph("Drafting notes", ss["Heading3"]))
        E += [Paragraph("&bull; " + escape(w), small) for w in doc.warnings]
    pdf.build(E, onFirstPage=marks, onLaterPages=marks)
    return path
