"""Shared python-docx helpers for JCE Shaw Road Rev B documents."""
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

GOLD = RGBColor(0xC9, 0xA2, 0x27)
BLACK = RGBColor(0x00, 0x00, 0x00)
GREY = RGBColor(0x55, 0x55, 0x55)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GOLD_HEX = "C9A227"
LIGHT_GOLD = "F6EDCF"


def shade(cell, hexcol):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcol)
    tcPr.append(shd)


def no_split_row(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def set_widths(t, widths):
    t.autofit = False
    grid = t._tbl.tblGrid
    for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
        if i < len(widths):
            gc.set(qn("w:w"), str(int(widths[i] * 567)))
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)


def bottom_rule(para, color=GOLD_HEX, sz="12"):
    pPr = para._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bot = OxmlElement("w:bottom")
    bot.set(qn("w:val"), "single")
    bot.set(qn("w:sz"), sz)
    bot.set(qn("w:space"), "1")
    bot.set(qn("w:color"), color)
    bdr.append(bot)
    pPr.append(bdr)


class Doc:
    def __init__(self, footer_text):
        self.d = Document()
        sec = self.d.sections[0]
        sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
        sec.left_margin = sec.right_margin = Cm(2.0)
        sec.top_margin, sec.bottom_margin = Cm(1.5), Cm(1.8)
        st = self.d.styles["Normal"]
        st.font.name = "Calibri"
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        st.font.size = Pt(10.5)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.line_spacing = 1.1
        self.sec_no = 0
        self.sub_no = 0
        self._footer(footer_text)

    # ---------- text ----------
    def _runs(self, para, text, bold=False, size=None, color=None, italic=False):
        for part in re.split(r"(\*\*[^*]+\*\*)", text):
            if not part:
                continue
            b = bold
            if part.startswith("**") and part.endswith("**"):
                part, b = part[2:-2], True
            r = para.add_run(part)
            r.bold = b
            r.italic = italic
            if size:
                r.font.size = Pt(size)
            if color is not None:
                r.font.color.rgb = color

    def p(self, text="", bold=False, size=None, color=None, align=None, after=None,
          before=None, italic=False, keep=False):
        para = self.d.add_paragraph()
        self._runs(para, text, bold=bold, size=size, color=color, italic=italic)
        if align is not None:
            para.alignment = align
        if after is not None:
            para.paragraph_format.space_after = Pt(after)
        if before is not None:
            para.paragraph_format.space_before = Pt(before)
        if keep:
            para.paragraph_format.keep_with_next = True
        return para

    def h(self, text):
        self.sec_no += 1
        self.sub_no = 0
        para = self.p(f"{self.sec_no}.0   {text.upper()}", bold=True, size=12.5, color=GOLD,
                      keep=True, before=12, after=4)
        bottom_rule(para, color="000000", sz="4")
        return para

    def sub(self, text):
        return self.p(text, bold=True, size=11, color=BLACK, keep=True, before=6, after=3)

    def bullets(self, items, style="List Bullet"):
        for it in items:
            para = self.d.add_paragraph(style=style)
            self._runs(para, it)
            para.paragraph_format.space_after = Pt(2)

    def numbered(self, items, prefix=None):
        for i, it in enumerate(items, 1):
            para = self.d.add_paragraph()
            para.paragraph_format.left_indent = Cm(0.9)
            para.paragraph_format.first_line_indent = Cm(-0.9)
            para.paragraph_format.space_after = Pt(3)
            label = f"{prefix}{i}" if prefix else f"{i}."
            self._runs(para, f"{label}\t{it}")
            para.paragraph_format.tab_stops.add_tab_stop(Cm(0.9))

    # ---------- tables ----------
    def table(self, header, rows, widths, total_rows=(), font=9.5, align_right_cols=()):
        t = self.d.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, htxt in enumerate(header):
            c = t.rows[0].cells[i]
            c.text = ""
            self._runs(c.paragraphs[0], htxt, bold=True, size=font, color=GOLD)
            shade(c, "000000")
        repeat_header(t.rows[0])
        n = len(rows)
        for ri, row in enumerate(rows):
            cells = t.add_row().cells
            is_total = (ri in total_rows) or (ri - n in total_rows)
            for i, val in enumerate(row):
                cells[i].text = ""
                para = cells[i].paragraphs[0]
                self._runs(para, str(val), bold=is_total, size=font)
                para.paragraph_format.space_after = Pt(1)
                if i in align_right_cols:
                    para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                if is_total:
                    shade(cells[i], LIGHT_GOLD)
        for ri, row in enumerate(t.rows):
            no_split_row(row)
            if ri < len(t.rows) - 1 and len(t.rows) <= 16:
                for c in row.cells:
                    for para in c.paragraphs:
                        para.paragraph_format.keep_with_next = True
        set_widths(t, widths)
        sp = self.d.add_paragraph()
        sp.paragraph_format.space_after = Pt(2)
        return t

    def kv_table(self, pairs, widths=(3.4, 5.1, 3.4, 5.1), font=9.5):
        t = self.d.add_table(rows=0, cols=4)
        t.style = "Table Grid"
        for a, b, c, d in pairs:
            cells = t.add_row().cells
            for i, (val, lab) in enumerate([(a, 1), (b, 0), (c, 1), (d, 0)]):
                cells[i].text = ""
                self._runs(cells[i].paragraphs[0], val, bold=bool(lab), size=font)
                cells[i].paragraphs[0].paragraph_format.space_after = Pt(1)
                if lab:
                    shade(cells[i], LIGHT_GOLD)
        for row in t.rows:
            no_split_row(row)
        set_widths(t, list(widths))
        self.d.add_paragraph().paragraph_format.space_after = Pt(0)

    # ---------- page furniture ----------
    def letterhead(self, doc_type, title, subtitle, our_ref, date, to_lines):
        self.p("JCE INFRASTRUCTURE PTY LTD", bold=True, size=17, color=BLACK, after=0)
        self.p("CIVIL, HYDRAULICS AND INFRASTRUCTURE ENGINEERING", bold=True, size=9, color=GOLD,
               after=1)
        self.p("ABN 29 699 165 663   |   0477 207 027   |   jayden@jceinfrastructure.com.au",
               size=9, color=GREY, after=0)
        para = self.p("756 Quia Road, Gunnedah NSW 2380   |   Servicing NSW and QLD", size=9,
                      color=GREY, after=8)
        bottom_rule(para, sz="18")
        self.p(f"**Our ref:** {our_ref}", size=10, after=0)
        self.p(f"**Date:** {date}", size=10, after=6)
        for i, line in enumerate(to_lines):
            self.p(line, size=10, after=0 if i < len(to_lines) - 1 else 8)
        self.p(doc_type, bold=True, size=10, color=GOLD, after=1)
        self.p(title, bold=True, size=15, color=BLACK, after=1)
        self.p(subtitle, size=10.5, color=GREY, after=8)

    def _footer(self, text):
        sec = self.d.sections[0]
        fp = sec.footer.paragraphs[0]
        fp.text = ""
        r = fp.add_run(text + "   |   Page ")
        r.font.size = Pt(8)
        r.font.color.rgb = GREY
        self._field(fp, "PAGE")
        r = fp.add_run(" of ")
        r.font.size = Pt(8)
        r.font.color.rgb = GREY
        self._field(fp, "NUMPAGES")

    def _field(self, para, code):
        run = para.add_run()
        run.font.size = Pt(8)
        run.font.color.rgb = GREY
        f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
        it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = code
        f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "end")
        run._r.append(f1); run._r.append(it); run._r.append(f2)

    def page_break(self):
        self.d.add_page_break()

    def figure(self, path, width_cm=17.0, caption=None):
        para = self.d.add_paragraph()
        para.paragraph_format.keep_with_next = True
        para.add_run().add_picture(path, width=Cm(width_cm))
        if caption:
            self.p(caption, italic=True, size=9, color=GREY)

    def signature_block(self, left_lines, right_lines, widths=(8.5, 8.5)):
        """Offer and acceptance in one unsplittable table row."""
        t = self.d.add_table(rows=1, cols=2)
        t.style = "Table Grid"
        for ci, lines in enumerate((left_lines, right_lines)):
            cell = t.rows[0].cells[ci]
            cell.text = ""
            first = True
            for ln in lines:
                para = cell.paragraphs[0] if first else cell.add_paragraph()
                first = False
                self._runs(para, ln, size=9.5)
                para.paragraph_format.keep_with_next = True
                para.paragraph_format.keep_together = True
                para.paragraph_format.space_after = Pt(2)
        no_split_row(t.rows[0])
        set_widths(t, list(widths))
        return t

    def save(self, path):
        self.d.save(path)


def money(v):
    return f"$ {v:,.2f}"
